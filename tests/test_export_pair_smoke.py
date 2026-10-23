from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from validate_export_pair import validate_pair

def test_generated_pptx_pdf_pair(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    from reportlab.pdfgen.canvas import Canvas

    pptx = tmp_path / "presentation.pptx"
    pdf = tmp_path / "presentation.pdf"
    deck = Presentation()
    deck.slide_width = Inches(13.333333)
    deck.slide_height = Inches(7.5)
    for i in range(3):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(9), Inches(1)).text = f"Slide {i+1}"
    deck.save(str(pptx))
    pdf_canvas = Canvas(str(pdf), pagesize=(960, 540))
    for i in range(3):
        pdf_canvas.drawString(50, 480, f"Slide {i+1}")
        pdf_canvas.showPage()
    pdf_canvas.save()
    assert validate_pair(pptx, pdf, 3) == []
    assert any("PDF page count" in e for e in validate_pair(pptx, pdf, 2))

def test_reject_missing_files(tmp_path):
    assert validate_pair(tmp_path/"absent.pptx",tmp_path/"absent.pdf",3)
