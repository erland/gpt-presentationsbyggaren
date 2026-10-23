#!/usr/bin/env python3
"""Verifiera levererade PPTX/PDF-artefakter, inte assistentens beteende."""
from __future__ import annotations
import argparse
from pathlib import Path
import sys

def validate_pair(pptx_path: Path, pdf_path: Path, expected_slides: int) -> list[str]:
    errors: list[str] = []
    try:
        from pptx import Presentation
        import fitz
    except ImportError as exc:
        return [f"Missing export validation dependency: {exc}"]
    root = Path(__file__).resolve().parent
    sys.path.insert(0, str(root))
    from validate_pptx import validate_pptx

    if not pptx_path.is_file() or not pdf_path.is_file():
        return ["Both PPTX and PDF must exist"]
    try:
        pptx_result = validate_pptx(str(pptx_path))
        errors.extend(str(e) for e in pptx_result['errors'])
    except Exception as exc:
        errors.append(f"PPTX integrity check failed: {exc}")
    try:
        presentation = Presentation(str(pptx_path))
        if len(presentation.slides) != expected_slides:
            errors.append(f"PPTX slide count {len(presentation.slides)} != {expected_slides}")
        ratio = presentation.slide_width / presentation.slide_height
        if abs(ratio - 16 / 9) > 0.02:
            errors.append(f"PPTX aspect ratio is {ratio:.3f}, expected 16:9")
    except Exception as exc:
        errors.append(f"PPTX could not be opened: {exc}")
    try:
        with fitz.open(str(pdf_path)) as doc:
            if len(doc) != expected_slides:
                errors.append(f"PDF page count {len(doc)} != {expected_slides}")
            if not doc or any(p.rect.width <= 0 or p.rect.height <= 0 for p in doc):
                errors.append("PDF has invalid page geometry")
            elif any(abs(p.rect.width / p.rect.height - 16 / 9) > 0.02 for p in doc):
                errors.append("PDF page aspect ratio is not 16:9")
    except Exception as exc:
        errors.append(f"PDF could not be opened: {exc}")
    return errors

def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx", type=Path)
    ap.add_argument("pdf", type=Path)
    ap.add_argument("--slides", required=True, type=int)
    args = ap.parse_args()
    errors = validate_pair(args.pptx, args.pdf, args.slides)
    for error in errors:
        print("FAIL:", error)
    print("ARTIFACT SMOKE:", "FAIL" if errors else "PASS")
    return 1 if errors else 0

if __name__ == "__main__":
    raise SystemExit(main())
