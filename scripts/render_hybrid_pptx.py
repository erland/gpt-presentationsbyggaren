#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

SLIDE_WIDTH_IN = 13.333333
SLIDE_HEIGHT_IN = 7.5

STYLE_PRESETS: dict[str, dict[str, Any]] = {
    "title-large": {"font_name": "Aptos Display", "font_size_pt": 28, "bold": True, "color": "#1F2937"},
    "subtitle-medium": {"font_name": "Aptos", "font_size_pt": 18, "bold": False, "color": "#374151"},
    "body": {"font_name": "Aptos", "font_size_pt": 18, "bold": False, "color": "#1F2937"},
    "label-large": {"font_name": "Aptos", "font_size_pt": 20, "bold": True, "color": "#1F2937"},
    "footnote": {"font_name": "Aptos", "font_size_pt": 10, "bold": False, "color": "#4B5563"},
}


def _pct(value: float, total_inches: float):
    return Inches(total_inches * value / 100.0)


def _rgb(hex_value: str) -> RGBColor:
    value = hex_value.strip().lstrip("#")
    if len(value) != 6:
        raise ValueError(f"Expected 6-digit RGB hex color, got: {hex_value}")
    return RGBColor.from_string(value.upper())



def _section(block: str, heading: str) -> str:
    start = block.find(heading)
    if start < 0:
        return ""
    start += len(heading)
    next_heading = re.search(r"^\*\*[^\n]+\*\*\s*$", block[start:], re.M)
    end = start + next_heading.start() if next_heading else len(block)
    return block[start:end].strip()


def _parse_keyed_bullets(section: str) -> dict[str, str]:
    result: dict[str, str] = {}
    for line in section.splitlines():
        match = re.match(r"^-\s*([^:]+):\s*(.+?)\s*$", line)
        if match:
            result[match.group(1).strip()] = match.group(2).strip()
    return result


def _parse_rendering_status(plan_text: str) -> dict[str, str]:
    start = plan_text.find("## Rendering status")
    end = plan_text.find("## Slides", start + 1)
    if start < 0 or end < 0:
        raise ValueError("Presentation plan is missing Rendering status")
    block = plan_text[start:end]
    return {
        slide_id: status
        for slide_id, status in re.findall(
            r"^- Slide\s+(\d{2,3}):\s*([a-z-]+)\s*$", block, re.M
        )
    }


def assert_packaging_ready(plan_path: str | Path) -> dict[str, str]:
    """Require every image/hybrid slide to be explicitly approved before packaging."""
    plan_path = Path(plan_path)
    text = plan_path.read_text(encoding="utf-8")
    statuses = _parse_rendering_status(text)

    blockers: list[str] = []
    for match in re.finditer(r"^### Slide\s+(\d{2,3})\s+—\s+.+$", text, re.M):
        slide_id = match.group(1)
        next_match = re.search(
            r"^### Slide\s+\d{2,3}\s+—\s+.+$",
            text[match.end():],
            re.M,
        )
        end = match.end() + next_match.start() if next_match else len(text)
        block = text[match.start():end]
        mode_match = re.search(r"\*\*Render mode:\*\*\s*(\S+)", block)
        if not mode_match or mode_match.group(1) not in {"image-slide", "hybrid-slide"}:
            continue
        status = statuses.get(slide_id)
        if status != "approved":
            blockers.append(f"Slide {slide_id}: {status or 'missing-status'}")

    if blockers:
        raise ValueError(
            "Presentation is not ready for PPTX packaging; "
            "image-slide and hybrid-slide must be approved: "
            + ", ".join(blockers)
        )
    return statuses


def _parse_layout_value(value: str) -> dict[str, Any]:
    parts = [part.strip() for part in value.split(",") if part.strip()]
    fields: dict[str, str] = {}
    for part in parts:
        if "=" not in part:
            continue
        key, raw = part.split("=", 1)
        fields[key.strip()] = raw.strip()

    def percent(name: str) -> float:
        raw = fields.get(name)
        if raw is None or not raw.endswith("%"):
            raise ValueError(f"Text layout requires {name}=NN%: {value}")
        return float(raw[:-1])

    return {
        "box": {
            "x_pct": percent("x"),
            "y_pct": percent("y"),
            "w_pct": percent("width"),
            "h_pct": percent("height"),
        },
        "style": fields.get("style", "body"),
        "align": fields.get("align", "left"),
    }


def parse_presentation_plan(plan_path: str | Path, assets_dir: str | Path) -> dict[str, Any]:
    """Project image-slide and hybrid-slide entries from presentation-plan.md."""
    plan_path = Path(plan_path)
    assets_dir = Path(assets_dir)
    text = plan_path.read_text(encoding="utf-8")
    statuses = _parse_rendering_status(text)

    matches = list(re.finditer(r"^### Slide\s+(\d{2,3})\s+—\s+(.+)$", text, re.M))
    slides: list[dict[str, Any]] = []

    for idx, match in enumerate(matches):
        start = match.start()
        end = matches[idx + 1].start() if idx + 1 < len(matches) else len(text)
        block = text[start:end]

        mode_match = re.search(r"\*\*Render mode:\*\*\s*(\S+)", block)
        if not mode_match:
            continue
        render_mode = mode_match.group(1)
        if render_mode not in {"image-slide", "hybrid-slide"}:
            continue

        slide_id = match.group(1)
        text_items: list[dict[str, Any]] = []

        if render_mode == "hybrid-slide":
            visible = _parse_keyed_bullets(_section(block, "**Visible text**"))
            layout = _parse_keyed_bullets(_section(block, "**Text layout**"))

            missing_layout = sorted(set(visible) - set(layout))
            if missing_layout:
                raise ValueError(
                    f"Slide {slide_id}: missing Text layout entries for: {', '.join(missing_layout)}"
                )

            for key, copy in visible.items():
                parsed = _parse_layout_value(layout[key])
                preset = dict(STYLE_PRESETS.get(parsed.pop("style"), STYLE_PRESETS["body"]))
                preset.update(parsed)
                preset["text"] = copy
                text_items.append(preset)
        background = None
        for candidate in (
            assets_dir / f"slide-{slide_id}.png",
            assets_dir / f"slide-{int(slide_id)}.png",
            assets_dir / f"slide-{slide_id}.jpg",
            assets_dir / f"slide-{int(slide_id)}.jpg",
        ):
            if candidate.is_file():
                background = str(candidate)
                break
        if background is None:
            raise FileNotFoundError(
                f"Slide {slide_id}: expected background asset in {assets_dir} "
                f"(for example slide-{slide_id}.png)"
            )

        slides.append(
            {
                "slide_id": slide_id,
                "title": match.group(2).strip(),
                "render_mode": render_mode,
                "rendering_status": statuses.get(slide_id),
                "background": background,
                "text": text_items,
            }
        )

    if not slides:
        raise ValueError("Presentation plan contains no renderable image-slide or hybrid-slide entries")
    return {"slides": slides}


def render_plan_hybrid_pptx(
    plan_path: str | Path, assets_dir: str | Path, output_path: str | Path
) -> Path:
    return render_hybrid_pptx(parse_presentation_plan(plan_path, assets_dir), output_path)


def render_plan_pptx(
    plan_path: str | Path, assets_dir: str | Path, output_path: str | Path
) -> Path:
    """Package a mixed visual-first deck after all image/hybrid slides are approved."""
    assert_packaging_ready(plan_path)
    return render_hybrid_pptx(parse_presentation_plan(plan_path, assets_dir), output_path)


def render_hybrid_pptx(spec: dict[str, Any], output_path: str | Path) -> Path:
    """Render image-backed slides with editable native PowerPoint text overlays."""
    output_path = Path(output_path)
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_WIDTH_IN)
    prs.slide_height = Inches(SLIDE_HEIGHT_IN)

    # Remove the default slide if a producer/library ever adds one.
    while prs.slides:
        r_id = prs.slides._sldIdLst[-1].rId
        prs.part.drop_rel(r_id)
        del prs.slides._sldIdLst[-1]

    blank_layout = prs.slide_layouts[6]

    for slide_spec in spec.get("slides", []):
        slide = prs.slides.add_slide(blank_layout)

        background = slide_spec.get("background")
        if background:
            background_path = Path(background)
            if not background_path.is_file():
                raise FileNotFoundError(background_path)
            slide.shapes.add_picture(
                str(background_path),
                0,
                0,
                width=prs.slide_width,
                height=prs.slide_height,
            )

        for item in slide_spec.get("text", []):
            box = item["box"]
            shape = slide.shapes.add_textbox(
                _pct(float(box["x_pct"]), SLIDE_WIDTH_IN),
                _pct(float(box["y_pct"]), SLIDE_HEIGHT_IN),
                _pct(float(box["w_pct"]), SLIDE_WIDTH_IN),
                _pct(float(box["h_pct"]), SLIDE_HEIGHT_IN),
            )
            frame = shape.text_frame
            frame.clear()
            frame.word_wrap = True

            paragraph = frame.paragraphs[0]
            paragraph.text = str(item["text"])

            align = str(item.get("align", "left")).lower()
            from pptx.enum.text import PP_ALIGN
            paragraph.alignment = {
                "left": PP_ALIGN.LEFT,
                "center": PP_ALIGN.CENTER,
                "right": PP_ALIGN.RIGHT,
            }.get(align, PP_ALIGN.LEFT)

            font = paragraph.runs[0].font
            font.name = item.get("font_name", "Aptos")
            font.size = Pt(float(item.get("font_size_pt", 28)))
            font.bold = bool(item.get("bold", False))
            font.color.rgb = _rgb(item.get("color", "#000000"))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    return output_path


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Package visual-first PPTX slides from image assets, with editable text overlays on hybrid slides."
    )
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument("--spec", help="JSON spec containing slides, backgrounds and text boxes")
    source.add_argument("--plan", help="Canonical presentation-plan.md")
    parser.add_argument("--assets-dir", help="Directory containing slide-NN.png/jpg background assets")
    parser.add_argument("--output", required=True, help="Output .pptx path")
    args = parser.parse_args()

    if args.plan:
        if not args.assets_dir:
            parser.error("--assets-dir is required with --plan")
        render_plan_pptx(args.plan, args.assets_dir, args.output)
    else:
        spec_path = Path(args.spec)
        spec = json.loads(spec_path.read_text(encoding="utf-8"))
        for slide in spec.get("slides", []):
            background = slide.get("background")
            if background and not Path(background).is_absolute():
                slide["background"] = str((spec_path.parent / background).resolve())
        render_hybrid_pptx(spec, args.output)

    print(f"HYBRID PPTX: {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
